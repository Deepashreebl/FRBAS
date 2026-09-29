from pathlib import Path
import sqlite3

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment
from openpyxl.drawing.image import Image as ExcelImage


# ============================================================
# PROJECT PATHS
# ============================================================

BASE_DIR = Path(__file__).resolve().parent.parent

DATABASE_PATH = BASE_DIR / "database" / "attendance.db"

EXCEL_DIR = BASE_DIR / "excel_reports"
EXCEL_DIR.mkdir(parents=True, exist_ok=True)

EXCEL_PATH = EXCEL_DIR / "attendance_database.xlsx"


# ============================================================
# DATABASE CONNECTION
# ============================================================

def get_connection():
    connection = sqlite3.connect(DATABASE_PATH)
    connection.row_factory = sqlite3.Row
    return connection


# ============================================================
# MAKE HEADER
# ============================================================

def format_header(worksheet):

    for cell in worksheet[1]:

        cell.font = Font(bold=True)

        cell.alignment = Alignment(
            horizontal="center",
            vertical="center"
        )


# ============================================================
# RESOLVE PHOTO PATH
# ============================================================

def resolve_photo_path(photo_path):

    if not photo_path:
        return None

    photo_file = Path(photo_path)

    if not photo_file.is_absolute():
        photo_file = BASE_DIR / photo_file

    if photo_file.exists():
        return photo_file

    return None


# ============================================================
# INSERT PHOTO INTO EXCEL
# ============================================================

def insert_photo(
    worksheet,
    photo_path,
    cell_reference,
    width=120,
    height=100
):

    photo_file = resolve_photo_path(photo_path)

    if not photo_file:
        return False

    if photo_file.suffix.lower() not in [
        ".jpg",
        ".jpeg",
        ".png"
    ]:
        return False

    try:

        excel_image = ExcelImage(
            str(photo_file)
        )

        excel_image.width = width
        excel_image.height = height

        worksheet.add_image(
            excel_image,
            cell_reference
        )

        return True

    except Exception as error:

        print(
            f"Warning: Could not insert photo "
            f"{photo_file}: {error}"
        )

        return False


# ============================================================
# EXPORT NORMAL TABLE
# ============================================================

def export_table(
    connection,
    workbook,
    sheet_name,
    table_name
):

    worksheet = workbook.create_sheet(sheet_name)

    cursor = connection.cursor()

    cursor.execute(
        f"PRAGMA table_info({table_name})"
    )

    columns = cursor.fetchall()

    if not columns:

        worksheet.append(["No data"])

        return

    column_names = [
        column["name"]
        for column in columns
    ]

    # --------------------------------------------------------
    # Header
    # --------------------------------------------------------

    worksheet.append(column_names)

    # --------------------------------------------------------
    # Data
    # --------------------------------------------------------

    cursor.execute(
        f"SELECT * FROM {table_name}"
    )

    rows = cursor.fetchall()

    for row in rows:

        worksheet.append(
            [
                row[column]
                for column in column_names
            ]
        )

    format_header(worksheet)

    worksheet.freeze_panes = "A2"

    # --------------------------------------------------------
    # Column widths
    # --------------------------------------------------------

    for column_cells in worksheet.columns:

        maximum_length = 0

        for cell in column_cells:

            if cell.value is not None:

                maximum_length = max(
                    maximum_length,
                    len(str(cell.value))
                )

        column_letter = (
            column_cells[0].column_letter
        )

        worksheet.column_dimensions[
            column_letter
        ].width = min(
            maximum_length + 3,
            40
        )


# ============================================================
# EXPORT EMPLOYEES WITH PHOTOS
# ============================================================

def export_employees(
    connection,
    workbook
):

    worksheet = workbook.create_sheet(
        "Employees"
    )

    cursor = connection.cursor()

    cursor.execute("""
        SELECT
            id,
            employee_id,
            name,
            email,
            phone,
            department,
            designation,
            salary,
            photo_path,
            created_at,
            is_active
        FROM employees
        ORDER BY id
    """)

    employees = cursor.fetchall()

    # --------------------------------------------------------
    # Headers
    # --------------------------------------------------------

    headers = [
        "ID",
        "Employee ID",
        "Name",
        "Email",
        "Phone",
        "Department",
        "Designation",
        "Salary",
        "Photo",
        "Open Photo",
        "Created At",
        "Active"
    ]

    worksheet.append(headers)

    format_header(worksheet)

    worksheet.freeze_panes = "A2"

    # --------------------------------------------------------
    # Employee rows
    # --------------------------------------------------------

    for row_number, employee in enumerate(
        employees,
        start=2
    ):

        # ----------------------------------------------------
        # Basic information
        # ----------------------------------------------------

        worksheet.cell(
            row=row_number,
            column=1,
            value=employee["id"]
        )

        worksheet.cell(
            row=row_number,
            column=2,
            value=employee["employee_id"]
        )

        worksheet.cell(
            row=row_number,
            column=3,
            value=employee["name"]
        )

        worksheet.cell(
            row=row_number,
            column=4,
            value=employee["email"]
        )

        worksheet.cell(
            row=row_number,
            column=5,
            value=employee["phone"]
        )

        worksheet.cell(
            row=row_number,
            column=6,
            value=employee["department"]
        )

        worksheet.cell(
            row=row_number,
            column=7,
            value=employee["designation"]
        )

        worksheet.cell(
            row=row_number,
            column=8,
            value=employee["salary"]
        )

        # ----------------------------------------------------
        # Employee Photo
        # ----------------------------------------------------

        photo_path = employee["photo_path"]

        photo_inserted = insert_photo(
            worksheet,
            photo_path,
            f"I{row_number}",
            width=100,
            height=100
        )

        open_photo_cell = worksheet.cell(
            row=row_number,
            column=10
        )

        photo_file = resolve_photo_path(
            photo_path
        )

        if photo_inserted:

            open_photo_cell.value = "Open Photo"

            open_photo_cell.hyperlink = (
                photo_file.as_uri()
            )

            open_photo_cell.font = Font(
                color="0000FF",
                underline="single"
            )

            worksheet.row_dimensions[
                row_number
            ].height = 80

        else:

            worksheet.cell(
                row=row_number,
                column=9,
                value="Photo not found"
            )

            open_photo_cell.value = (
                "Not available"
            )

        # ----------------------------------------------------
        # Remaining information
        # ----------------------------------------------------

        worksheet.cell(
            row=row_number,
            column=11,
            value=employee["created_at"]
        )

        worksheet.cell(
            row=row_number,
            column=12,
            value=(
                "Yes"
                if employee["is_active"]
                else "No"
            )
        )

    # --------------------------------------------------------
    # Column widths
    # --------------------------------------------------------

    widths = {
        "A": 8,
        "B": 15,
        "C": 25,
        "D": 30,
        "E": 18,
        "F": 20,
        "G": 20,
        "H": 15,
        "I": 18,
        "J": 18,
        "K": 25,
        "L": 12
    }

    for column, width in widths.items():

        worksheet.column_dimensions[
            column
        ].width = width

    # --------------------------------------------------------
    # Alignment
    # --------------------------------------------------------

    for row in worksheet.iter_rows(
        min_row=2,
        max_row=worksheet.max_row
    ):

        for column_number in [
            1,
            2,
            5,
            8,
            9,
            10,
            12
        ]:

            row[
                column_number - 1
            ].alignment = Alignment(
                horizontal="center",
                vertical="center"
            )


# ============================================================
# EXPORT ATTENDANCE WITH ACTUAL PHOTOS
# ============================================================

def export_attendance(
    connection,
    workbook
):

    worksheet = workbook.create_sheet(
        "Attendance"
    )

    cursor = connection.cursor()

    cursor.execute("""
        SELECT
            id,
            employee_id,
            date,
            check_in,
            check_out,
            working_hours,
            status,
            check_in_photo,
            check_out_photo
        FROM attendance
        ORDER BY date DESC, id DESC
    """)

    attendance_rows = cursor.fetchall()

    # --------------------------------------------------------
    # Headers
    # --------------------------------------------------------

    headers = [
        "ID",
        "Employee ID",
        "Date",
        "Check In",
        "Check Out",
        "Working Hours",
        "Status",
        "Check-In Photo",
        "Open Check-In Photo",
        "Check-Out Photo",
        "Open Check-Out Photo"
    ]

    worksheet.append(headers)

    format_header(worksheet)

    worksheet.freeze_panes = "A2"

    # --------------------------------------------------------
    # Attendance rows
    # --------------------------------------------------------

    for row_number, attendance in enumerate(
        attendance_rows,
        start=2
    ):

        # ----------------------------------------------------
        # Basic attendance information
        # ----------------------------------------------------

        worksheet.cell(
            row=row_number,
            column=1,
            value=attendance["id"]
        )

        worksheet.cell(
            row=row_number,
            column=2,
            value=attendance["employee_id"]
        )

        worksheet.cell(
            row=row_number,
            column=3,
            value=attendance["date"]
        )

        worksheet.cell(
            row=row_number,
            column=4,
            value=attendance["check_in"]
        )

        worksheet.cell(
            row=row_number,
            column=5,
            value=attendance["check_out"]
        )

        worksheet.cell(
            row=row_number,
            column=6,
            value=attendance["working_hours"]
        )

        worksheet.cell(
            row=row_number,
            column=7,
            value=attendance["status"]
        )

        # ----------------------------------------------------
        # CHECK-IN PHOTO
        # ----------------------------------------------------

        check_in_photo = attendance[
            "check_in_photo"
        ]

        check_in_file = resolve_photo_path(
            check_in_photo
        )

        check_in_inserted = insert_photo(
            worksheet,
            check_in_photo,
            f"H{row_number}",
            width=120,
            height=100
        )

        open_check_in_cell = worksheet.cell(
            row=row_number,
            column=9
        )

        if check_in_inserted:

            open_check_in_cell.value = (
                "Open Check-In Photo"
            )

            open_check_in_cell.hyperlink = (
                check_in_file.as_uri()
            )

            open_check_in_cell.font = Font(
                color="0000FF",
                underline="single"
            )

        else:

            worksheet.cell(
                row=row_number,
                column=8,
                value="Photo not found"
            )

            open_check_in_cell.value = (
                "Not available"
            )

        # ----------------------------------------------------
        # CHECK-OUT PHOTO
        # ----------------------------------------------------

        check_out_photo = attendance[
            "check_out_photo"
        ]

        check_out_file = resolve_photo_path(
            check_out_photo
        )

        check_out_inserted = insert_photo(
            worksheet,
            check_out_photo,
            f"J{row_number}",
            width=120,
            height=100
        )

        open_check_out_cell = worksheet.cell(
            row=row_number,
            column=11
        )

        if check_out_inserted:

            open_check_out_cell.value = (
                "Open Check-Out Photo"
            )

            open_check_out_cell.hyperlink = (
                check_out_file.as_uri()
            )

            open_check_out_cell.font = Font(
                color="0000FF",
                underline="single"
            )

        else:

            worksheet.cell(
                row=row_number,
                column=10,
                value="Photo not found"
            )

            open_check_out_cell.value = (
                "Not available"
            )

        # ----------------------------------------------------
        # Row height for photos
        # ----------------------------------------------------

        if (
            check_in_inserted
            or check_out_inserted
        ):

            worksheet.row_dimensions[
                row_number
            ].height = 85

    # --------------------------------------------------------
    # Column widths
    # --------------------------------------------------------

    widths = {
        "A": 8,
        "B": 18,
        "C": 15,
        "D": 15,
        "E": 15,
        "F": 18,
        "G": 15,
        "H": 22,
        "I": 25,
        "J": 22,
        "K": 25
    }

    for column, width in widths.items():

        worksheet.column_dimensions[
            column
        ].width = width

    # --------------------------------------------------------
    # Alignment
    # --------------------------------------------------------

    for row in worksheet.iter_rows(
        min_row=2,
        max_row=worksheet.max_row
    ):

        for cell in row:

            cell.alignment = Alignment(
                horizontal="center",
                vertical="center"
            )

# ============================================================
# MAIN EXPORT FUNCTION
# ============================================================

def export_database_to_excel():

    if not DATABASE_PATH.exists():

        print()
        print("ERROR: Database not found.")

        print(
            f"Expected location: {DATABASE_PATH}"
        )

        return

    connection = get_connection()

    # --------------------------------------------------------
    # Create workbook
    # --------------------------------------------------------

    workbook = Workbook()

    default_sheet = workbook.active

    workbook.remove(default_sheet)

    # --------------------------------------------------------
    # Employees + Photos
    # --------------------------------------------------------

    export_employees(
        connection,
        workbook
    )

    # --------------------------------------------------------
    # Attendance + Check-In/Check-Out Photos
    # --------------------------------------------------------

    export_attendance(
        connection,
        workbook
    )

    # --------------------------------------------------------
    # Other database tables
    # --------------------------------------------------------

    export_table(
        connection,
        workbook,
        "Leaves",
        "leaves"
    )

    export_table(
        connection,
        workbook,
        "Payroll",
        "payroll"
    )

    export_table(
        connection,
        workbook,
        "Audit Logs",
        "audit_logs"
    )

    connection.close()

    # --------------------------------------------------------
    # Save Excel
    # --------------------------------------------------------

    workbook.save(
        EXCEL_PATH
    )

    print()
    print("=" * 70)
    print(
        "       EXCEL DATABASE UPDATED SUCCESSFULLY"
    )
    print("=" * 70)
    print()

    print("Excel file:")
    print(EXCEL_PATH)

    print()

    print("Sheets created:")
    print("1. Employees")
    print("2. Attendance")
    print("3. Leaves")
    print("4. Payroll")
    print("5. Audit Logs")

    print()

    print(
        "Employee photos are embedded in the Employees sheet."
    )

    print(
        "Check-in and check-out photos are embedded "
        "in the Attendance sheet."
    )

    print(
        "Photo file paths are NOT displayed "
        "in the Attendance sheet."
    )

    print("=" * 70)


# ============================================================
# PROGRAM START
# ============================================================

if __name__ == "__main__":

    export_database_to_excel()